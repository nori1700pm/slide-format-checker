"""Mechanical python-pptx helpers for the slide-format-checker skill.

These only apply formatting. Classification (cover, divider, checklist, header vs
body), dates, overlap and alignment are decided by SKILL.md, not here.
"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

TITLE_FONT = "HDharmony B"
BODY_FONT = "Hyundai Sans Text KR"
NAVY = RGBColor(0x00, 0x20, 0x60)

# Exact positions in EMU (1 inch = 914400). Inches() truncates, so use constants.
TITLE_POS = (274320, 228600)               # 0.3", 0.25"
CHECKLIST_TITLE_POS = (155448, 271681)     # template slide 2
DESCRIPTION_POS = (457200, 868680)         # 0.5", 0.95"

# Children that may follow a:latin / a:ea inside a:rPr (schema order).
_AFTER_LATIN = ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")
_AFTER_EA = _AFTER_LATIN[1:]


def is_hidden(slide):
    return slide._element.get("show") == "0"


def iter_text_frames(shapes):
    """Yield every text frame in shapes, recursing into groups and table cells."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_frames(shape.shapes)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
        elif shape.has_text_frame:
            yield shape.text_frame


def iter_rprs(paragraph):
    """Yield the a:rPr of every run and field in a paragraph (creating it if absent)."""
    for child in paragraph._p:
        if child.tag in (qn("a:r"), qn("a:fld")):
            rPr = child.find(qn("a:rPr"))
            if rPr is None:
                rPr = OxmlElement("a:rPr")
                child.insert(0, rPr)
            yield rPr


def _get_or_add(rPr, tag, successors):
    el = rPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        nxt = next((c for c in rPr if c.tag in {qn(s) for s in successors}), None)
        if nxt is not None:
            nxt.addprevious(el)
        else:
            rPr.append(el)
    return el


def set_typeface(rPr, face):
    """Set both the Latin and East Asian typeface, keeping schema order."""
    _get_or_add(rPr, "a:latin", _AFTER_LATIN).set("typeface", face)
    _get_or_add(rPr, "a:ea", _AFTER_EA).set("typeface", face)


def format_paragraph_runs(paragraph, face=None, size=None, rgb=None, bold=None):
    """Apply typeface / size (pt) / color / bold to every run and field in a paragraph."""
    for rPr in iter_rprs(paragraph):
        if face is not None:
            set_typeface(rPr, face)
        if size is not None:
            rPr.set("sz", str(int(Pt(size).pt * 100)))
        if bold is not None:
            rPr.set("b", "1" if bold else "0")
        if rgb is not None:
            for old in rPr.findall(qn("a:solidFill")):
                rPr.remove(old)
            fill = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", str(rgb))
            fill.append(clr)
            # fill sits before latin/ea/cs/sym/hlink in a:rPr
            nxt = next((c for c in rPr if c.tag in {qn(s) for s in ("a:latin",) + _AFTER_LATIN}), None)
            if nxt is not None:
                nxt.addprevious(fill)
            else:
                rPr.append(fill)


def format_title(shape, checklist=False):
    """Position + HDharmony B / 24 pt / #002060 / bold off. Caller has done the overlap check."""
    shape.left, shape.top = CHECKLIST_TITLE_POS if checklist else TITLE_POS
    for p in shape.text_frame.paragraphs:
        format_paragraph_runs(p, TITLE_FONT, 24, NAVY, False)


def format_brief_description(shape):
    """Position + HDharmony B / 18 pt / #002060 / bold off. Caller has done the overlap check."""
    shape.left, shape.top = DESCRIPTION_POS
    for p in shape.text_frame.paragraphs:
        format_paragraph_runs(p, TITLE_FONT, 18, NAVY, False)


def at_position(shape, pos):
    """Exact EMU match, no tolerance."""
    return (shape.left, shape.top) == tuple(pos)


def has_text(paragraph):
    """Non-empty paragraph. Digits and symbols count; only whitespace does not."""
    return bool("".join(t.text or "" for t in paragraph._p.iter(qn("a:t"))).strip())


def turn_off_latin_wrap(paragraph):
    """Set latinLnBrk="0". Returns True only when the value actually changed."""
    if not has_text(paragraph):
        return False
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.get("latinLnBrk") == "0":
        return False
    pPr.set("latinLnBrk", "0")
    return True


def clamp_body_size(size_pt):
    """Body size rule: above 12 -> 12, below 11 -> 11, otherwise unchanged (None = no change)."""
    if size_pt is None:
        return None
    if size_pt > 12:
        return 12
    if size_pt < 11:
        return 11
    return None


def font_installed(name):
    """Windows only: is the font in the font registry (system or per-user)?"""
    import winreg
    key_path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
    pattern = re.compile(r"^" + re.escape(name) + r"( \((TrueType|OpenType)\))?$", re.I)
    for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            key = winreg.OpenKey(hive, key_path)
        except OSError:
            continue
        i = 0
        while True:
            try:
                value_name = winreg.EnumValue(key, i)[0]
            except OSError:
                break
            if pattern.match(value_name):
                return True
            i += 1
    return False
